"""Export KPIs in the exact row/column layout of the `O&M KPI` sheet.

Same labels in column B (verbatim, including the original spelling so donor
reports still line up), months in C:N, the annual figure in O. Values are
numbers, not formulas: every ratio is computed by the KPI engine. Empty cell =
no data (never 0 by default).
"""
import csv
import io

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

from ops.models import FailureCause as FC
from ops.models import NRWCause as NC

from .service import MONTHS_FR

MONTH_HEADERS = ["Janvier", "Fevrier", "Mars", "Avril", "Mai", "Juin", "Juillet", "Aout", "Septembre", "Octobre", "Novembre", "Décembre"]

# (row, label, metric, format, annual_metric)
#   metric None + format "header" -> section header row with month names
#   format: "n" number, "pct" percentage, "usd" currency, "h" hours, "m3" cubic metres
H = "header"
LAYOUT = [
    (4, "Disponibilité du réseau", None, H, None),
    (5, "Heure d'Arrêt du système dû à la panne des pompes", "downtime_pump", "h", "downtime_pump"),
    (6, "Heure d'Arrêt du système dû à la rupture des conduites", "downtime_pipe", "h", "downtime_pipe"),
    (7, "Heure d'arrêt du système dû à la coupe d'électricité", "downtime_power", "h", "downtime_power"),
    (8, "Heure d'arrêt du système dû à une maintenance programmée", "downtime_planned", "h", "downtime_planned"),
    (9, "Total heure d'arrêt", "downtime_total", "h", "downtime_total"),
    (10, "Heure réelle de fonctionnement du réseau pour 24hr/jour", "operating_hours", "h", "operating_hours"),
    (11, "Taux de disponibilité du réseau", "availability", "pct", "availability"),
    (12, "Nombre des pannes dans le réseau", None, H, None),
    (13, "Nombre/mois", "incidents_reported", "n", "incidents_reported"),
    (14, "Nombre total des pannes/an", None, "n", "incidents_reported"),
    (15, "Nombre des pannes réparées", "incidents_closed", "n", "incidents_closed"),
    (16, "Taux de reparation des pannes", "repair_rate", "pct", "repair_rate"),
    (17, "Analyse de cause racine (RCA) des pannes", None, H, "Taux annuel"),
    (18, "Vandalisme et vole", f"rca_{FC.VANDALISM}", "n", f"rca_share_{FC.VANDALISM}"),
    (19, "Surpression", f"rca_{FC.OVERPRESSURE}", "n", f"rca_share_{FC.OVERPRESSURE}"),
    (20, "Tuyau mal enfuie", f"rca_{FC.SHALLOW_PIPE}", "n", f"rca_share_{FC.SHALLOW_PIPE}"),
    (21, "Connexion illégale", f"rca_{FC.ILLEGAL_CONNECTION}", "n", f"rca_share_{FC.ILLEGAL_CONNECTION}"),
    (22, "Mauvaise manipulation des accessoires le technicien", f"rca_{FC.MISHANDLING}", "n", f"rca_share_{FC.MISHANDLING}"),
    (23, "Mauvaise qualité du tuyau", f"rca_{FC.POOR_PIPE_QUALITY}", "n", f"rca_share_{FC.POOR_PIPE_QUALITY}"),
    (24, "Mauvais remblai", f"rca_{FC.POOR_BACKFILL}", "n", f"rca_share_{FC.POOR_BACKFILL}"),
    (25, "Mouvement de terrain", f"rca_{FC.GROUND_MOVEMENT}", "n", f"rca_share_{FC.GROUND_MOVEMENT}"),
    (26, "Mauvaise installation", f"rca_{FC.POOR_INSTALLATION}", "n", f"rca_share_{FC.POOR_INSTALLATION}"),
    (27, "Coup de bélier", f"rca_{FC.WATER_HAMMER}", "n", f"rca_share_{FC.WATER_HAMMER}"),
    (28, "Connexion défectueuse", f"rca_{FC.FAULTY_CONNECTION}", "n", f"rca_share_{FC.FAULTY_CONNECTION}"),
    (29, "Rendement du réseau", None, H, None),
    (30, "Eau introduite", "volume_introduced", "m3", "volume_introduced"),
    (31, "Eau comptabilisée", "volume_billed", "m3", "volume_billed"),
    (32, "Rendement", "efficiency", "pct", "efficiency"),
    (33, "Eau non facturée (NRW) en m3", "nrw_m3", "m3", "nrw_m3"),
    (34, "Nombre de fois il y a eu perte d'eau", "nrw_events", "n", "nrw_events"),
    (35, "Analyse de cause racine NRW", None, H, "Taux annuel"),
    (36, "Fuites physiques sur les conduites", f"nrw_{NC.PIPE_LEAK}", "n", f"nrw_share_{NC.PIPE_LEAK}"),
    (37, "Ruptures de conduites", f"nrw_{NC.PIPE_BURST}", "n", f"nrw_share_{NC.PIPE_BURST}"),
    (38, "Débordements de réservoirs-faute de trop pleins", f"nrw_{NC.RESERVOIR_OVERFLOW}", "n", f"nrw_share_{NC.RESERVOIR_OVERFLOW}"),
    (39, "Purges du réseau", f"nrw_{NC.NETWORK_FLUSH}", "n", f"nrw_share_{NC.NETWORK_FLUSH}"),
    (40, "Vidanges", f"nrw_{NC.DRAINING}", "n", f"nrw_share_{NC.DRAINING}"),
    (41, "Eau utilisée pour le nettoyage des réservoirs", f"nrw_{NC.RESERVOIR_CLEANING}", "n", f"nrw_share_{NC.RESERVOIR_CLEANING}"),
    (42, "Eau utilisée pour la lutte contre les incendies", f"nrw_{NC.FIREFIGHTING}", "n", f"nrw_share_{NC.FIREFIGHTING}"),
    (43, "Branchements illégaux", f"nrw_{NC.ILLEGAL_BRANCH}", "n", f"nrw_share_{NC.ILLEGAL_BRANCH}"),
    (44, "Compteurs défectueux", f"nrw_{NC.FAULTY_METER}", "n", f"nrw_share_{NC.FAULTY_METER}"),
    (45, "Compteurs mal calibrés", f"nrw_{NC.MISCALIBRATED_METER}", "n", f"nrw_share_{NC.MISCALIBRATED_METER}"),
    (46, "Connexions illégales", f"nrw_{NC.ILLEGAL_CONNECTION}", "n", f"nrw_share_{NC.ILLEGAL_CONNECTION}"),
    (47, "Erreurs de relevé", f"nrw_{NC.READING_ERROR}", "n", f"nrw_share_{NC.READING_ERROR}"),
    (48, "Volumes consommés mais non facturés", f"nrw_{NC.UNBILLED_CONSUMPTION}", "n", f"nrw_share_{NC.UNBILLED_CONSUMPTION}"),
    (49, "Taux de maintenance préventive", None, H, None),
    (50, "Nombre de maintenance au Captage", "pm_planned_CAPTAGE", "n", "pm_planned_CAPTAGE"),
    (51, "Nombre de maintenance Pompage", "pm_planned_POMPAGE", "n", "pm_planned_POMPAGE"),
    (52, "Nombre de maintenance Réservoir", "pm_planned_RESERVOIR", "n", "pm_planned_RESERVOIR"),
    (53, "Nombre de maintenance Réseau", "pm_planned_RESEAU", "n", "pm_planned_RESEAU"),
    (54, "Nombre de maintenance Bornes-fontaines", "pm_planned_BF", "n", "pm_planned_BF"),
    (55, "Total prévu pour tous les actifs", "pm_planned_total", "n", "pm_planned_total"),
    (56, "Réalisées", "pm_done_total", "n", "pm_done_total"),
    (57, "Taux", "pm_rate", "pct", "pm_rate"),
    (58, "Consommation énergétique", None, H, None),
    (59, "Consommation électrcité en kWh", "kwh", "n", "kwh"),
    (60, "Intensité kWh/m3", "kwh_per_m3", "r", "kwh_per_m3"),
    (61, "Consommation  carburant en litre", "fuel_l", "n", "fuel_l"),
    (62, "Intensité litre/m3", "fuel_l_per_m3", "r", "fuel_l_per_m3"),
    (63, "Coût Electricité (USD)", "cost_electricity", "usd", "cost_electricity"),
    (64, "Coût Carburant (USD)", "cost_fuel", "usd", "cost_fuel"),
    (65, "Estimation Kwh USD/m3", "electricity_cost_per_m3", "r", None),
    (66, "Estimation carburant USD/m3", "fuel_cost_per_m3", "r", None),
    (67, "Coût énergétique total (USD)", "energy_cost", "usd", "energy_cost"),
    (68, "Coût énergétique total (USD)/m³", "energy_cost_per_m3", "r", "energy_cost_per_m3"),
    (69, "Coût d'Exploitation et de maintenance", None, H, None),
    (70, "Budget prévu (USD)", "budget", "usd", "budget"),
    (71, "Maintenance Urgente (USD)", "cost_urgent", "usd", "cost_urgent"),
    (72, "Maintenance Corrective (USD)", "cost_corrective", "usd", "cost_corrective"),
    (73, "Maintenance Préventive (USD)", "cost_preventive", "usd", "cost_preventive"),
    (74, "Autres activités de support (USD)", "cost_support", "usd", "cost_support"),
    (75, "Total Dépense réelle (USD)", "actual_total", "usd", "actual_total"),
    (76, "Écart budgétaire", "budget_variance", "usd", "budget_variance"),
    (77, "Maintenance urgente %", "share_urgent", "pct", "share_urgent"),
    (78, "Maintenance corrective %", "share_corrective", "pct", "share_corrective"),
    (79, "Maintenance préventive %", "share_preventive", "pct", "share_preventive"),
    (80, "Autres activités de support %", "share_support", "pct", "share_support"),
    (81, "Qualité conforme", None, H, None),
    (82, "I. Nombre de mesure du chlore résiduel", "quality_field_total", "n", "quality_field_total"),
    (83, "Mesures conformes", "quality_field_compliant", "n", "quality_field_compliant"),
    (84, "Mesures non-conformes", "quality_field_noncompliant", "n", "quality_field_noncompliant"),
    (85, "II. Nombre de tests laboratoires périodiques", "quality_lab_total", "n", "quality_lab_total"),
    (86, "Mesures conformes", "quality_lab_compliant", "n", "quality_lab_compliant"),
    (87, "Mesures non-conformes", "quality_lab_noncompliant", "n", "quality_lab_noncompliant"),
    (88, "Taux de conformité %", "quality_rate", "pct", "quality_rate"),
]

FORMATS = {"n": "#,##0.##", "h": "#,##0.##", "m3": "#,##0.##", "usd": "#,##0.00", "pct": "0.0%", "r": "0.000"}


def _with_derived(values):
    v = dict(values)
    for p in ("quality_field", "quality_lab"):
        t, c = v.get(f"{p}_total"), v.get(f"{p}_compliant")
        v[f"{p}_noncompliant"] = (t - (c or 0)) if t is not None else None
    return v


def build_workbook(result, site_name):
    wb = Workbook()
    ws = wb.active
    ws.title = "O&M KPI"
    bold = Font(bold=True)
    head_fill = PatternFill("solid", fgColor="DDE6EE")
    thin = Side(style="thin", color="999999")
    ws["B2"] = f"INDICATEURS DE PERFORMANCE D'E&M DU SYSTEME (KPI)- Réseau {site_name.upper()} - YME JIBU "
    ws["B2"].font = Font(bold=True, size=13)
    ws["B3"] = "KPI"
    ws["B3"].font = bold
    ws.column_dimensions["B"].width = 58
    for c in range(3, 16):
        ws.column_dimensions[get_column_letter(c)].width = 12

    months = [_with_derived(m["values"]) for m in result["months"]]
    annual = _with_derived(result["annual"])
    for row, label, metric, fmt, annual_metric in LAYOUT:
        ws.cell(row, 2, label)
        if fmt == H:
            ws.cell(row, 2).font = bold
            for i, name in enumerate(MONTH_HEADERS):
                cell = ws.cell(row, 3 + i, name)
                cell.font = bold
                cell.fill = head_fill
            ws.cell(row, 15, annual_metric or "Année").font = bold
            ws.cell(row, 15).fill = head_fill
            ws.cell(row, 2).fill = head_fill
            continue
        if metric:
            for i, mv in enumerate(months):
                val = mv.get(metric)
                if val is not None:
                    cell = ws.cell(row, 3 + i, float(val))
                    cell.number_format = FORMATS[fmt]
        if annual_metric:
            val = annual.get(annual_metric)
            if val is not None:
                cell = ws.cell(row, 15, float(val))
                cell.number_format = FORMATS["pct"] if annual_metric.startswith(("rca_share", "nrw_share")) else FORMATS[fmt]
        for c in range(2, 16):
            ws.cell(row, c).border = Border(bottom=thin)
        ws.cell(row, 2).alignment = Alignment(wrap_text=True)
    ws.freeze_panes = "C5"

    notes = wb.create_sheet("Lisez-moi")
    lines = [
        "Export généré par la plateforme E&M Yme Jibu.",
        f"Site : {site_name} — Année {result['year']} — Données au {result['today']}.",
        "Même disposition que la feuille « O&M KPI » : lignes et libellés identiques, mois en C:N, année en O.",
        "Toutes les valeurs sont calculées à partir des fiches (relevés, pannes, ordres de travail, dépenses).",
        "Pour les mois antérieurs à l'application, les totaux mensuels validés de l'ancien fichier Excel sont utilisés ;",
        "les ratios sont toujours recalculés (aucune formule Excel n'est reprise).",
        "Cellule vide = pas de donnée. Une valeur n'est jamais remplacée par 0 ou 100 % faute de données.",
        "Taux annuels = ratio des sommes (jamais une moyenne de ratios).",
        "« Total Dépense réelle » inclut les autres activités de support (ligne 74), omises dans l'ancien fichier.",
        "",
        "Mois | Statut | Jours avec relevés de pompage | Sources",
    ]
    for m in result["months"]:
        srcs = sorted(set(m["sources"].values()))
        lines.append(f"{MONTHS_FR[m['month'] - 1]} | {m['status']} | {m['coverage']['pump_reading_days']}/{m['coverage']['days']} | {', '.join(srcs) or '—'}")
    for i, line in enumerate(lines, start=1):
        notes.cell(i, 1, line)
    notes.column_dimensions["A"].width = 120
    return wb


def workbook_bytes(result, site_name):
    buf = io.BytesIO()
    build_workbook(result, site_name).save(buf)
    return buf.getvalue()


def csv_text(result):
    buf = io.StringIO()
    w = csv.writer(buf, delimiter=";")
    w.writerow(["ligne", "indicateur"] + MONTH_HEADERS + ["annee"])
    months = [_with_derived(m["values"]) for m in result["months"]]
    annual = _with_derived(result["annual"])
    for row, label, metric, fmt, annual_metric in LAYOUT:
        if fmt == H:
            continue
        vals = [months[i].get(metric) if metric else None for i in range(12)]
        a = annual.get(annual_metric) if annual_metric else None
        w.writerow([row, label] + ["" if v is None else f"{float(v):.4f}".rstrip("0").rstrip(".") for v in vals]
                   + ["" if a is None else f"{float(a):.4f}".rstrip("0").rstrip(".")])
    return buf.getvalue()
