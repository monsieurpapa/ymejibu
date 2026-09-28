"""The KPI export keeps the exact layout of the `O&M KPI` sheet."""
import io

import openpyxl
import pytest
from django.conf import settings

from .conftest import TODAY, needs_workbooks
from importer.loaders import import_all
from importer.xl import Workbooks
from kpi.export import LAYOUT
from kpi.service import compute_year

pytestmark = pytest.mark.django_db


@needs_workbooks
def test_labels_match_original_sheet():
    original = Workbooks(settings.WORKBOOK_DIR).f("kpi", "O&M KPI")
    for row, label, *_ in LAYOUT:
        assert (original.cell(row, 2).value or "").strip() == label.strip(), f"ligne {row}"


@needs_workbooks
def test_xlsx_export_values(db):
    from rest_framework.test import APIClient

    from .conftest import make_user
    from core.models import Role

    wb = Workbooks(settings.WORKBOOK_DIR)
    site, _ = import_all(wb, 2026, TODAY)
    client = APIClient()
    client.force_authenticate(make_user(site, "bailleur_go", Role.FUNDER))
    resp = client.get("/api/kpi/export.xlsx?year=2026")
    assert resp.status_code == 200
    out = openpyxl.load_workbook(io.BytesIO(resp.content))["O&M KPI"]
    original = wb.f("kpi", "O&M KPI")
    for row in range(4, 89):
        assert out.cell(row, 2).value == original.cell(row, 2).value
    assert out["C4"].value == "Janvier" and out["N4"].value == "Décembre"
    assert out["C30"].value == 600 and out["C31"].value == 495
    assert out["C33"].value == 105 and out["D33"].value == 105  # NRW every month
    assert out["C62"].value == pytest.approx(0.375)  # 225 L / 600 m³
    assert out["E10"].value == 730  # 744 − 14
    assert out["C16"].value is None  # no invented repair rate
    assert out["L30"].value is None  # October: future, empty
    csv = client.get("/api/kpi/export.csv?year=2026").content.decode("utf-8-sig")
    assert csv.splitlines()[0].startswith("ligne;indicateur;Janvier")


# --- PDF report -------------------------------------------------------------------------------
def test_pdf_report_annual_and_monthly(client_for, site):
    import io as _io

    from pypdf import PdfReader

    c = client_for("resp")
    r = c.get("/api/kpi/report.pdf?year=2026")
    assert r.status_code == 200 and r["Content-Type"] == "application/pdf"
    assert r.content[:4] == b"%PDF" and "ymejibu_indicateurs_TS_2026.pdf" in r["Content-Disposition"]
    text = "".join(p.extract_text() for p in PdfReader(_io.BytesIO(r.content)).pages)
    assert "Rapport des indicateurs E&M 2026" in text and "Disponibilité du réseau" in text and "Tableau mensuel 2026" in text

    r = c.get("/api/kpi/report.pdf?year=2026&month=8")
    assert r.status_code == 200 and "2026-08" in r["Content-Disposition"]
    text = "".join(p.extract_text() for p in PdfReader(_io.BytesIO(r.content)).pages)
    assert "Rapport mensuel — Août 2026" in text and "Tendance sur l'année" in text

    assert c.get("/api/kpi/report.pdf?year=2026&month=13").status_code == 400


def test_pdf_report_follows_dashboard_roles(client_for, site):
    assert client_for("tech").get("/api/kpi/report.pdf").status_code == 403
    assert client_for("funder").get("/api/kpi/report.pdf").status_code == 200


def test_kpi_api_serves_catalog(client_for, site):
    data = client_for("resp").get("/api/kpi/?year=2026").json()
    keys = [k["key"] for k in data["catalog"]["kpis"]]
    assert "availability" in keys and data["catalog"]["breakdowns"][0]["items"]
    assert all(k["key"] in data["months"][0]["values"] for k in data["catalog"]["kpis"])
